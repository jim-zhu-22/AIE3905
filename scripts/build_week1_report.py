"""Build the week-one Word report from actual saved experiment results."""

import argparse
import json
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


LABELS = ("paper", "plastic", "metal", "wood", "fabric")
ZH = {"paper": "纸质", "plastic": "塑料", "metal": "金属", "wood": "木材", "fabric": "织物"}


def build(root: Path, run: Path, output: Path):
    metrics = json.loads((run / "evaluation/metrics.json").read_text(encoding="utf-8"))
    history = json.loads((run / "history.json").read_text(encoding="utf-8"))
    config = json.loads((run / "config.json").read_text(encoding="utf-8"))
    best_index = max(range(len(history)), key=lambda i: history[i]["validation"]["macro_f1"])
    total_minutes = sum(row["elapsed_seconds"] for row in history) / 60
    head_epochs = sum(row["stage"] == "head" for row in history)
    full_epochs = sum(row["stage"] == "full" for row in history)
    gpu_metrics = json.loads((root / "runs/week1_gpu_probe/evaluation/metrics.json").read_text(encoding="utf-8"))
    f1 = metrics["classification"]["macro avg"]["f1-score"]
    matrix = metrics["confusion_matrix"]
    pairs = sorted(((matrix[i][j], LABELS[i], LABELS[j]) for i in range(5) for j in range(5) if i != j), reverse=True)
    weak = min(LABELS, key=lambda label: metrics["classification"][label]["recall"])
    figures = root / "docs/week1_results"
    doc = Document()
    for border in list(doc.styles.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(0.65)
    sec.left_margin = sec.right_margin = Inches(0.7)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(6)
    doc.styles["Normal"].font.size = Pt(11)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.15
    doc.styles["Title"].font.size = Pt(23)
    doc.styles["Heading 1"].font.size = Pt(15)
    doc.styles["Heading 2"].font.size = Pt(12)
    doc.core_properties.title = "仓储机器人表面材质识别第一周周报"
    doc.core_properties.author = "MirrorSpace × AIE3905 项目组"

    def paragraph(text):
        return doc.add_paragraph(text)

    def heading(text):
        doc.add_heading(text, 1)

    def table(headers, rows, widths):
        t = doc.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for c, width in zip(t.columns, widths):
            c.width = Inches(width)
        for c, value in zip(t.rows[0].cells, headers):
            c.text = str(value)
        for values in rows:
            for c, value in zip(t.add_row().cells, values):
                c.text = str(value)
        for ri, row in enumerate(t.rows):
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            if ri == 0:
                row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            for ci, cell in enumerate(row.cells):
                cell.width = Inches(widths[ci])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                pr = cell._tc.get_or_add_tcPr()
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "244A67" if ri == 0 else ("F1F5F8" if ri % 2 else "FFFFFF"))
                pr.append(shade)
                borders = OxmlElement("w:tcBorders")
                for edge in ("top", "left", "bottom", "right"):
                    e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), "single")
                    e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "D9D9D9"); borders.append(e)
                pr.append(borders)
                margins = OxmlElement("w:tcMar")
                for edge in ("top", "left", "bottom", "right"):
                    e = OxmlElement(f"w:{edge}"); e.set(qn("w:w"), "90"); e.set(qn("w:type"), "dxa"); margins.append(e)
                pr.append(margins)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.space_before = Pt(2)
                    for r in p.runs:
                        r.font.size = Pt(10.5)
                        if ri == 0:
                            r.font.bold = True; r.font.color.rgb = RGBColor(255, 255, 255)
        doc.add_paragraph().paragraph_format.space_after = Pt(1)
        return t

    def figure(name, width):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(figures / name), width=Inches(width))

    doc.add_paragraph("仓储机器人表面材质识别\n第一周周报", "Title")
    paragraph(f"MirrorSpace × AIE3905 课程项目  |  报告日期 {datetime.now():%Y年%m月%d日}")
    paragraph(f"本周完成公开数据整理、ResNet-18 小样本与完整数据训练、验证评估和 ROI 展示流程。完整数据验证准确率为 {metrics['accuracy']:.2%}，Macro-F1 为 {f1:.4f}。公开数据线已具备第一周交付条件；现场数据线仍等待企业图像与标注导出样例。")
    heading("第一周指标与完成状态")
    table(("验收项", "实际完成情况", "状态"), [
        ("公开数据", "五类共 12,500 张，保留官方划分；图片可解码且跨划分无字节级重复", "完成"),
        ("模型基线", f"完整训练集 10,625 张，验证集 625 张；完成 {len(history)} 轮训练并保存最佳权重", "完成"),
        ("评估证据", "逐类指标、混淆矩阵、训练曲线及逐图预测记录", "完成"),
        ("展示与规范", "上传已裁剪 ROI；统一五类标签；清单模板和初版标注规范", "完成"),
        ("现场数据接入", "待收到现有标注格式、ROI 和 SKU 字段；新 SKU 测试尚未实施", "待数据"),
        ("GPU 环境", "Python 3.12 CUDA 环境已跑通 200/50 张小样本训练；完整基线使用 CPU", "完成验证"),
    ], (1.1, 4.8, 0.9))
    heading("数据与运行环境")
    paragraph("公开数据为 MINC-2500 split 1，选取 paper、plastic、metal、wood、fabric。训练集 10,625 张，验证集 625 张，保留测试集 1,250 张；五类在每个集合中均衡。本周未使用测试集进行选模或评分。")
    paragraph(f"完整基线使用 Python 3.11、PyTorch {config['torch']}、CPU。另在已有 Python 3.12.12 / PyTorch 2.9.1+cu128 / CUDA 12.8 环境，使用 RTX 4050 6 GB 跑通小规模 GPU 训练。显卡连续测速较慢，32 张批量前反向约 14 秒，系统报告低频与限速；原因尚未确认。")

    doc.add_page_break()
    heading("完整数据训练与验证结果")
    paragraph(f"采用 ImageNet 预训练 ResNet-18，输入 224 × 224，替换五类分类头。先训练分类头 {head_epochs} 轮，学习率 {config['head_lr']}；再从分类头最佳权重初始化，微调整个网络 {full_epochs} 轮，学习率 {config['full_lr']}。批量大小 {config['batch_size']}，两个阶段随机种子均为 {config['seed']}，优化器 AdamW。分类头阶段冻结骨干及其 BatchNorm 统计。")
    paragraph(f"按验证集 Macro-F1 选择最佳权重，最佳为总第 {best_index + 1} 轮。累计训练与每轮验证耗时约 {total_minutes:.1f} 分钟，不含环境排查和最终图表制作。")
    table(("指标", "结果", "样本范围"), [
        ("Accuracy", f"{metrics['accuracy']:.2%}", "625 张公开验证图"),
        ("Macro-F1", f"{f1:.4f}", "五类等权平均"),
        ("测试集结果", "未评估", "保留 1,250 张"),
        ("新 SKU 泛化", "未评估", "需现场数据和 SKU 字段"),
        ("GPU 小样本 Accuracy", f"{gpu_metrics['accuracy']:.2%}", "200 张训练、50 张验证，独立演示"),
    ], (1.5, 1.5, 3.8))
    figure("training_curves.png", 6.6)
    paragraph("曲线用于检查训练趋势与验证表现。验证集参与选模，因此以上结果是开发阶段指标，不能作为独立最终测试成绩。此前 250 张小样本试验的验证准确率为 52%，仅用于确认流程；由于验证样本不同，不直接计算两者的提升幅度。")

    doc.add_page_break()
    heading("逐类表现与混淆分析")
    table(("材质", "Precision", "Recall", "F1", "样本数"), [
        (f"{ZH[label]} {label}", f"{metrics['classification'][label]['precision']:.3f}",
         f"{metrics['classification'][label]['recall']:.3f}", f"{metrics['classification'][label]['f1-score']:.3f}", "125")
        for label in LABELS
    ], (2.1, 1.2, 1.2, 1.2, 1.1))
    figure("confusion_matrix.png", 5.5)
    paragraph("混淆矩阵的行为真实类别，列为预测类别。最常见的错误包括：" + "；".join(f"{ZH[a]}误判为{ZH[b]} {n} 张" for n, a, b in pairs[:3]) + "。")
    paragraph(f"当前召回率最低的类别为{ZH[weak]}。下一阶段应检查这类错误中的反光、纹理、覆膜和混合表面，并在统一验证协议下比较局部 ROI 与完整商品图像。公开验证成绩本身不能证明模型识别了真实物理材质。")

    doc.add_page_break()
    heading("预测样例与本周限制")
    figure("prediction_examples.png", 6.8)
    paragraph("上排选取每类一个预测正确的样本，下排选取模型分数较高的错误样本，便于分析失败原因；它们并非随机抽样。图中分数为 softmax 输出，未做概率校准，也未实现完整未知材质拒识。")
    paragraph("Demo 输入为已裁剪的单材质接触面 RGB ROI，未完成自动物体定位。MINC 图像仍可能包含背景与多种表面，需结合局部 ROI 分析错误。MINC 不含现场 SKU，尚不能验证新商品泛化；织物标签也不能推断为纯棉。双线模型切换和真空吸盘验证尚未完成。")
    heading("第二周计划")
    for text in (
        "检查显卡低频与限速状态，恢复后复核完整数据 GPU 训练；固定数据划分并记录设备与训练预算。",
        "收到现场数据和标注导出样例后，适配 ROI、SKU、采集批次及人工确认状态，固定最终测试 SKU，建立现场基线。",
        "完善模型来源与版本展示，分析错误样例；按计划冻结公开 POC 配置后再进行保留测试，区分公开与现场成绩。",
    ):
        doc.add_paragraph(text, "List Bullet")
    heading("交付与复现位置")
    paragraph(f"项目目录 E:\\AIE3905。训练权重、配置与历史在 {run.relative_to(root).as_posix()}；逐图预测在该目录 evaluation；图表在 docs/week1_results。运行方法见 README.md，标注边界见 docs/annotation_guideline.md。数据与模型文件保留在本地，不进入代码仓库。")
    paragraph("数据来源 https://huggingface.co/datasets/mcimpoi/minc-2500_split_1；版本 5ac92ad1af8ef34f9ecd8d54d60a73c0289fe0c3。周报指标来自本地保存的训练历史和验证结果。")
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    summary = [
        "# 仓储机器人表面材质识别第一周周报", "",
        f"报告日期：{datetime.now():%Y-%m-%d}。公开数据线完成第一周基线；企业数据线等待图像与标注。", "",
        f"完整验证集 Accuracy **{metrics['accuracy']:.2%}**，Macro-F1 **{f1:.4f}**；验证集参与选模，不是独立测试成绩。", "",
        "## 完成情况", "",
        "- 已下载五类 MINC 全集 12,500 张：训练 10,625、验证 625、测试 1,250；保留官方划分。",
        "- 图片解码与清单检查通过，跨集合无字节级重复。",
        f"- 完成分类头 {head_epochs} 轮、全网络 {full_epochs} 轮，最佳为总第 {best_index+1} 轮。",
        f"- 实际训练设备：{config['device']}，PyTorch {config['torch']}；训练及每轮验证合计 {total_minutes:.1f} 分钟。",
        "- 已准备模型权重、ROI Demo、逐图预测、训练曲线、混淆矩阵及标注规范。", "",
        "## 逐类验证指标", "", "| 材质 | Precision | Recall | F1 | 数量 |", "| --- | ---: | ---: | ---: | ---: |",
    ]
    for label in LABELS:
        row = metrics["classification"][label]
        summary.append(f"| {ZH[label]} | {row['precision']:.3f} | {row['recall']:.3f} | {row['f1-score']:.3f} | 125 |")
    summary += ["", "## 每轮记录", "", "| 总轮次 | 阶段 | 训练准确率 | 验证准确率 | 验证 Macro-F1 | 秒 |", "| ---: | --- | ---: | ---: | ---: | ---: |"]
    for i, row in enumerate(history, 1):
        summary.append(f"| {i} | {row['stage']} | {row['train']['accuracy']:.2%} | {row['validation']['accuracy']:.2%} | {row['validation']['macro_f1']:.4f} | {row['elapsed_seconds']:.1f} |")
    summary += ["", "## 结果图", "", "![训练曲线](week1_results/training_curves.png)", "", "![混淆矩阵](week1_results/confusion_matrix.png)", "", "![预测样例](week1_results/prediction_examples.png)", "",
                "样例为每类一个正确样本和高分错误样本，不是随机抽样；softmax 分数尚未做概率校准。", "",
                "## 未完成项与下一周", "",
                f"- Python 3.12 CUDA 环境的小样本训练已完成，验证准确率 {gpu_metrics['accuracy']:.2%}；完整基线使用 CPU，数据规模不同不可直接比较。",
                "- GPU 初期同步异常、连续运算低效，原因未确认；恢复显卡状态后再完成全量 GPU 训练。见 gpu_environment.md。",
                "- 企业数据尚未提供，未完成现场基线、跨 SKU 测试或真空吸盘验证。",
                "- 下周固定 GPU 环境，适配企业标注与 ROI/SKU 字段，固定测试 SKU；冻结公开 POC 后再评估保留测试集。",
                "- 当前模型只识别已知五类，不包含完整 unknown 拒识能力。", "",
                f"权重、配置和实验历史：`{run.relative_to(root).as_posix()}`。Word 周报：`{output.relative_to(root).as_posix()}`。"]
    (root / "docs/week1_report.md").write_text("\n".join(summary) + "\n", encoding="utf-8")
    print(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--run", type=Path, default=Path("runs/week1_resnet18"))
    parser.add_argument("--output", type=Path, default=Path("deliverables/第一周周报.docx"))
    args = parser.parse_args()
    build(args.root.resolve(), (args.root / args.run).resolve(), (args.root / args.output).resolve())
